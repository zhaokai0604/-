"""上传后 analysis-live：真实 pipeline + done 落库。"""

from __future__ import annotations

from fastapi.testclient import TestClient


def _register(client: TestClient, username: str = "live_stream_user") -> None:
    response = client.post(
        "/api/auth/register",
        json={
            "username": username,
            "password": "Password1",
            "confirm_password": "Password1",
            "display_name": username,
        },
    )
    assert response.status_code == 200


def test_analyze_stream_then_live_persists(client: TestClient, tmp_path):
    _register(client)
    resume = tmp_path / "cv-live.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("王五")
    doc.add_paragraph("数据分析实习生")
    doc.add_paragraph("熟悉 Python SQL Excel，完成过数据清洗与可视化项目")
    doc.save(resume)

    create = client.post(
        "/api/resumes/analyze",
        files={"file": ("cv-live.docx", resume.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={
            "target_position": "数据分析师",
            "enable_ai": "false",
            "stream": "true",
        },
    )
    assert create.status_code == 200
    body = create.json()
    assert body.get("stream") is True
    assert body.get("status") == "processing"
    record_id = body["record_id"]

    status = client.get(f"/api/history/{record_id}/status")
    assert status.status_code == 200
    # 可能仍在 processing，也可能已被后台兜底跑完
    assert status.json()["status"] in {"processing", "success"}

    with client.stream("GET", f"/api/history/{record_id}/analysis-live") as stream:
        assert stream.status_code == 200
        text = "".join(stream.iter_text())
    assert "event: done" in text
    assert '"persisted": true' in text or "分析完成" in text or '"already_done": true' in text

    detail = client.get(f"/api/history/{record_id}")
    assert detail.status_code == 200
    payload = detail.json()
    assert payload["status"] == "success"
    assert payload["total_score"] > 0
    assert payload.get("rewrite_preview", {}).get("items") or payload.get("scores")


def test_live_stream_rejects_foreign_record(client: TestClient, tmp_path):
    _register(client, "live_owner")
    resume = tmp_path / "cv.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("测试")
    doc.add_paragraph("Python")
    doc.save(resume)
    create = client.post(
        "/api/resumes/analyze",
        files={"file": ("cv.docx", resume.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"enable_ai": "false", "stream": "true"},
    )
    record_id = create.json()["record_id"]

    other = TestClient(client.app)
    other.post(
        "/api/auth/register",
        json={
            "username": "live_other",
            "password": "Password1",
            "confirm_password": "Password1",
            "display_name": "other",
        },
    )
    denied = other.get(f"/api/history/{record_id}/analysis-live")
    assert denied.status_code == 404
