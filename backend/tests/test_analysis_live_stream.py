"""上传后真实 SSE 分析并落库。"""

from __future__ import annotations

import base64

from fastapi.testclient import TestClient


def _register(client: TestClient, username: str = "live_stream_user") -> None:
    response = client.post(
        "/api/auth/register",
        json={
            "username": username,
            "password": "Password1",
            "confirm_password": "Password1",
            "display_name": username,
        },
    )
    assert response.status_code == 200


def test_analyze_direct_stream_then_live_persist(client: TestClient, tmp_path):
    _register(client)
    resume = tmp_path / "cv-live.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("李四")
    doc.add_paragraph("数据分析实习生")
    doc.add_paragraph("熟悉 Python SQL Excel，完成过业务数据分析")
    doc.save(resume)

    created = client.post(
        "/api/resumes/analyze-direct",
        json={
            "filename": "cv-live.docx",
            "content_base64": base64.b64encode(resume.read_bytes()).decode("ascii"),
            "target_position": "数据分析师",
            "enable_ai": False,
            "stream": True,
        },
    )
    assert created.status_code == 200
    body = created.json()
    assert body.get("stream") is True
    assert body.get("status") == "processing"
    record_id = body["record_id"]

    # stream=true 时不应被后台立刻跑完
    status = client.get(f"/api/history/{record_id}/status")
    assert status.status_code == 200
    assert status.json()["status"] == "processing"

    with client.stream("GET", f"/api/history/{record_id}/analysis-live") as stream:
        assert stream.status_code == 200
        text = "".join(stream.iter_text())
    assert "event: done" in text
    assert "已落库" in text or "persisted" in text

    detail = client.get(f"/api/history/{record_id}")
    assert detail.status_code == 200
    payload = detail.json()
    assert payload["status"] == "success"
    assert payload["total_score"] > 0


def test_direct_target_payload_is_ignored_when_matching_disabled(client: TestClient, tmp_path):
    """岗位名和岗位库 ID 不能在关闭开关时污染通用分析。"""
    _register(client, "target_gate_user")
    from docx import Document

    resume = tmp_path / "new-media-resume.docx"
    doc = Document()
    doc.add_paragraph("求职者")
    doc.add_paragraph("新媒体专业")
    doc.add_paragraph("参与校园公众号内容编辑与活动宣传")
    doc.save(resume)

    created = client.post(
        "/api/resumes/analyze-direct",
        json={
            "filename": resume.name,
            "content_base64": base64.b64encode(resume.read_bytes()).decode("ascii"),
            "target_position": "嵌入式开发实习生",
            "job_description": "需要 C++、单片机和嵌入式开发经验",
            "job_profile_id": 999999,
            "target_match_enabled": False,
            "enable_ai": False,
            "stream": False,
        },
    )
    assert created.status_code == 200
    record_id = created.json()["record_id"]

    from app.api.platform import process_single_analysis

    process_single_analysis(record_id)
    detail = client.get(f"/api/history/{record_id}")
    assert detail.status_code == 200
    payload = detail.json()
    assert payload["target_position"] == ""
    assert payload["target_position_source"] == "generic"
    assert payload["match_result"]["target_position"] == ""
    assert payload["match_result"]["target_source"] == "generic"
    assert payload["job_market_match"]["recommendation_only"] is True
    assert payload["rewrite_preview"]["target_position"] in {"", "目标岗位"}
    assert all(
        position not in str(payload["rewrite_preview"])
        for position in ("嵌入式开发实习生", "数据分析师", "新媒体运营")
    )
    assert payload["template_recommendations"]["target_position"] == "通用"
