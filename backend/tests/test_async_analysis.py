"""Tests for async single analysis."""

import base64

from fastapi.testclient import TestClient


def _register(client: TestClient, username: str = "async_user") -> None:
    response = client.post(
        "/api/auth/register",
        json={"username": username, "password": "Password1", "confirm_password": "Password1", "display_name": username},
    )
    assert response.status_code == 200


def test_single_analysis_returns_processing_and_completes(client: TestClient, tmp_path):
    _register(client)
    record_id = _upload_sample_resume(client, tmp_path)

    detail = client.get(f"/api/history/{record_id}")
    assert detail.status_code == 200
    body = detail.json()
    assert body["status"] == "success"
    assert body["total_score"] > 0
    assert body["analysis_mode"] == "core"
    assert body["analysis_mode_label"] == "规则分析"
    assert body["ai_enhancement_status"] == "none"
    assert body.get("interview_prep", {}) == {}
    assert body.get("mock_interview", {}) == {}
    assert body.get("rewrite_preview", {}).get("items")
    assert "blocks" not in body

    interview = client.post(f"/api/history/{record_id}/interview-prep/refresh")
    assert interview.status_code == 200
    assert interview.json().get("interview_prep", {}).get("question_count", 0) > 0
    assert interview.json().get("mock_interview", {}).get("total_steps", 0) > 0

    rewrite = client.post(f"/api/history/{record_id}/rewrite-preview/refresh")
    assert rewrite.status_code == 200
    assert rewrite.json().get("rewrite_preview", {}).get("items")


def _upload_sample_resume(client: TestClient, tmp_path) -> int:
    resume = tmp_path / "cv.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("张三")
    doc.add_paragraph("数据分析师")
    doc.add_paragraph("Python SQL Excel 负责数据分析项目")
    doc.save(resume)

    with resume.open("rb") as handle:
        response = client.post(
            "/api/resumes/analyze",
            files={"file": ("cv.docx", handle, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"target_position": "数据分析师", "enable_ai": "false"},
        )
    assert response.status_code == 200
    body = response.json()
    assert body.get("async") is True
    return body["record_id"]


def test_retry_failed_single_analysis(client: TestClient, tmp_path):
    _register(client)
    record_id = _upload_sample_resume(client, tmp_path)

    from app.core.database import SessionLocal
    from app.models.entities import AnalysisRecord

    with SessionLocal() as db:
        record = db.query(AnalysisRecord).filter(AnalysisRecord.id == record_id).first()
        assert record is not None
        record.status = "failed"
        record.error_message = "模拟失败"
        db.commit()

    retry = client.post(f"/api/history/{record_id}/retry")
    assert retry.status_code == 200
    assert retry.json()["status"] == "processing"

    detail = client.get(f"/api/history/{record_id}")
    assert detail.status_code == 200
    assert detail.json()["status"] == "success"
    assert detail.json()["total_score"] > 0


def test_single_analysis_direct_upload_returns_processing_and_completes(client: TestClient, tmp_path):
    _register(client, username="async_direct_user")
    resume = tmp_path / "cv-direct.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("鏉庡洓")
    doc.add_paragraph("鍚庣寮€鍙戝疄涔犵敓")
    doc.add_paragraph("Python FastAPI SQL")
    doc.save(resume)

    response = client.post(
        "/api/resumes/analyze-direct",
        json={
            "filename": "cv-direct.docx",
            "content_base64": base64.b64encode(resume.read_bytes()).decode("ascii"),
            "target_position": "鍚庣寮€鍙戝疄涔犵敓",
            "enable_ai": False,
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body.get("async") is True

    detail = client.get(f"/api/history/{body['record_id']}")
    assert detail.status_code == 200
    assert detail.json()["status"] == "success"


def test_ai_enhancement_updates_success_record(client: TestClient, monkeypatch):
    from app.api.platform import process_ai_enhancement
    from app.core.database import SessionLocal
    from app.models.entities import AnalysisRecord, User
    from app.utils.json_tools import dumps, loads

    with SessionLocal() as db:
        user = db.query(User).filter(User.username == "guest").first()
        record = AnalysisRecord(
            user_id=user.id,
            original_filename="ai.docx",
            target_position="数据分析师",
            total_score=80,
            scores_json=dumps({"job_match": 80}),
            sections_json=dumps(
                {
                    "skills": ["Python SQL"],
                    "internship": ["负责日常数据分析工作"],
                    "_rewrite_preview": {"mode": "offline_star", "items": [{"section": "经历", "original": "负责日常数据分析工作", "suggested": "负责日常数据分析工作", "focus": "规则"}]},
                    "_structured_suggestions": [],
                    "_parse_quality": "high",
                    "_parse_warnings": [],
                    "_ai_requested": True,
                    "_ai_enhancement_status": "pending",
                    "_ai_enhancement_error": "",
                }
            ),
            diagnosis_json=dumps(["经历描述可以更量化"]),
            suggestions_json=dumps(["补充项目指标"]),
            match_result_json=dumps({"target_position": "数据分析师", "target_source": "manual", "missing_keywords": ["Excel"], "matched_keywords": ["Python"]}),
            analysis_mode="core",
            status="success",
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        record_id = record.id

    def fake_enhance(result, resume_text, enable_ai):
        result["diagnosis"] = ["AI 诊断"]
        result["suggestions"] = ["AI 建议"]
        result["structured_suggestions"] = [{"problem": "AI 问题", "evidence": "Python", "impact": "匹配弱", "direction": "补充", "example": "AI 示例"}]
        result["rewrite_preview"] = {"mode": "deepseek", "summary": "AI 改写", "items": [{"section": "经历", "original": "负责日常数据分析工作", "suggested": "使用 Python 和 SQL 输出 __ 份报告。", "focus": "AI"}]}
        return result, "deepseek"

    monkeypatch.setattr("app.services.platform_service.enhance_with_deepseek", fake_enhance)
    process_ai_enhancement(record_id)

    with SessionLocal() as db:
        record = db.query(AnalysisRecord).filter(AnalysisRecord.id == record_id).first()
        sections = loads(record.sections_json, {})
        assert record.status == "success"
        assert record.analysis_mode == "deepseek"
        assert sections["_ai_enhancement_status"] == "success"
        assert loads(record.diagnosis_json, [])[0] == "AI 诊断"
        assert sections["_rewrite_preview"]["mode"] == "deepseek"

        from app.api.platform import record_to_response

        payload = record_to_response(record, include_detail=True, db=db)
        assert payload["analysis_mode_label"] == "增强分析已完成"


def test_ai_enhancement_failure_returns_core_label(client: TestClient, monkeypatch):
    from app.api.platform import process_ai_enhancement, record_to_response
    from app.core.database import SessionLocal
    from app.models.entities import AnalysisRecord, User
    from app.utils.json_tools import dumps, loads

    with SessionLocal() as db:
        user = db.query(User).filter(User.username == "guest").first()
        record = AnalysisRecord(
            user_id=user.id,
            original_filename="ai-failed.docx",
            target_position="数据分析师",
            total_score=76,
            scores_json=dumps({"job_match": 76}),
            sections_json=dumps(
                {
                    "skills": ["Python SQL"],
                    "_rewrite_preview": {"mode": "offline_star", "items": []},
                    "_structured_suggestions": [],
                    "_ai_requested": True,
                    "_ai_enhancement_status": "pending",
                    "_ai_enhancement_error": "",
                }
            ),
            diagnosis_json=dumps(["规则诊断"]),
            suggestions_json=dumps(["规则建议"]),
            match_result_json=dumps({"target_position": "数据分析师"}),
            analysis_mode="core",
            status="success",
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        pending_payload = record_to_response(record, include_detail=True, db=db)
        assert pending_payload["analysis_mode_label"] == "增强分析进行中"
        record_id = record.id

    def fake_enhance(result, resume_text, enable_ai):
        result["ai_fallback_reason"] = "DeepSeek unavailable"
        return result, "offline_fallback"

    monkeypatch.setattr("app.services.platform_service.enhance_with_deepseek", fake_enhance)
    process_ai_enhancement(record_id)

    with SessionLocal() as db:
        record = db.query(AnalysisRecord).filter(AnalysisRecord.id == record_id).first()
        sections = loads(record.sections_json, {})
        assert record.status == "success"
        assert record.analysis_mode == "core"
        assert record.error_message == ""
        assert sections["_ai_enhancement_status"] == "failed"
        assert sections["_ai_enhancement_error"] == "DeepSeek unavailable"

        payload = record_to_response(record, include_detail=True, db=db)
        assert payload["analysis_mode_label"] == "规则分析"
        assert payload["ai_fallback_reason"] == ""
        assert payload["ai_enhancement_error"] == "DeepSeek unavailable"


def test_rewrite_report_download(client: TestClient, tmp_path):
    _register(client)
    record_id = _upload_sample_resume(client, tmp_path)
    refresh = client.post(f"/api/history/{record_id}/rewrite-preview/refresh")
    assert refresh.status_code == 200

    response = client.get(f"/api/history/{record_id}/rewrite-report")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in response.headers.get("content-type", "")
    assert len(response.content) > 1000


def test_report_download_is_generated_on_demand(client: TestClient, tmp_path):
    _register(client, "report_user")
    record_id = _upload_sample_resume(client, tmp_path)

    from app.core.database import SessionLocal
    from app.models.entities import Report

    with SessionLocal() as db:
        assert db.query(Report).filter(Report.analysis_record_id == record_id).count() == 0

    response = client.get(f"/api/history/{record_id}/report/download?format=docx")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in response.headers.get("content-type", "")
    assert len(response.content) > 1000

    with SessionLocal() as db:
        assert db.query(Report).filter(Report.analysis_record_id == record_id, Report.format == "docx").count() == 1
