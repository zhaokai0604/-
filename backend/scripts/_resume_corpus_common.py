"""真实简历流水线公共工具：枚举文件、抽正文、安全读写。

硬约束（三脚本共用）：
- 禁止调用大模型给匹配打 1–5 分
- 禁止把真实/脱敏简历自动写入 train_sbert 训练集
- 禁止把未脱敏原文写入可提交评测 JSON
"""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Any, Iterable

SUPPORTED_SUFFIXES = {".docx", ".pdf", ".txt", ".md"}
SKIP_NAME_PREFIXES = ("~$", ".")


def project_root() -> Path:
    return Path(__file__).resolve().parents[2]


def default_raw_dir() -> Path:
    return project_root() / "data" / "raw_real_resumes"


def default_anon_dir() -> Path:
    return project_root() / "data" / "anonymized_real_resumes"


def default_pools_dir() -> Path:
    return project_root() / "data" / "resume_pools"


def iter_resume_files(root: Path) -> list[Path]:
    if not root.exists():
        return []
    files: list[Path] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        name = path.name
        if name.startswith(SKIP_NAME_PREFIXES):
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        # 跳过明显非简历的说明文件
        if name.lower() in {"readme.md", "readme.txt", ".gitkeep"}:
            continue
        files.append(path)
    return files


def safe_stem(path: Path) -> str:
    stem = re.sub(r"[^\w\u4e00-\u9fa5\-]+", "_", path.stem).strip("_")
    return stem[:80] or "resume"


def extract_text(
    path: Path,
    *,
    enable_ocr: bool = False,
    ocr_engine: str = "paddleocr",
) -> tuple[str, list[str]]:
    """尽量抽正文；失败不抛到顶层，由调用方记入报告。"""
    warnings: list[str] = []
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md"}:
            raw = path.read_text(encoding="utf-8", errors="ignore")
            return raw, warnings
        if suffix == ".docx":
            return _extract_docx(path, warnings)
        if suffix == ".pdf":
            return _extract_pdf(path, warnings, enable_ocr=enable_ocr, ocr_engine=ocr_engine)
    except Exception as exc:  # noqa: BLE001 - 批处理要吞掉单文件异常
        warnings.append(f"提取失败: {type(exc).__name__}: {exc}")
        return "", warnings
    warnings.append(f"不支持的后缀: {suffix}")
    return "", warnings


def _extract_docx(path: Path, warnings: list[str]) -> tuple[str, list[str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx，请在 backend venv 中安装依赖") from exc

    doc = Document(str(path))
    chunks: list[str] = []
    seen: set[str] = set()

    def add(value: str) -> None:
        text = re.sub(r"\s+", " ", (value or "").replace("\u3000", " ")).strip()
        if text and text not in seen:
            seen.add(text)
            chunks.append(text)

    for paragraph in doc.paragraphs:
        add(paragraph.text)
    for table in doc.tables:
        _walk_table(table, add)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            add(paragraph.text)
        for paragraph in section.footer.paragraphs:
            add(paragraph.text)

    # 文本框等高信号 XML 碎片
    xml_bits: list[str] = []
    try:
        for element in doc.part.element.iter():
            if element.tag.endswith("}t") and element.text:
                add(element.text)
                xml_bits.append(element.text)
    except Exception:  # noqa: BLE001
        warnings.append("docx XML 文本框扫描跳过")

    text = "\n".join(chunks)
    if len(text.strip()) < 40:
        warnings.append("docx 抽到正文过短，可能是文本框/图片模板")
    if not text.strip() and xml_bits:
        text = "\n".join(xml_bits)
        warnings.append("回退使用 XML 碎片文本")
    return text, warnings


def _walk_table(table: Any, add) -> None:
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            cell_text = " ".join(
                re.sub(r"\s+", " ", p.text).strip() for p in cell.paragraphs if p.text and p.text.strip()
            )
            if cell_text:
                cells.append(cell_text)
            for nested in getattr(cell, "tables", []) or []:
                _walk_table(nested, add)
        if cells:
            add(" | ".join(cells))


def _extract_pdf(
    path: Path,
    warnings: list[str],
    *,
    enable_ocr: bool = False,
    ocr_engine: str = "paddleocr",
) -> tuple[str, list[str]]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("缺少 pdfplumber，请在 backend venv 中安装依赖") from exc

    chunks: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        page_count = len(pdf.pages)
        if page_count == 0:
            warnings.append("PDF 无页面")
        if page_count > 20:
            warnings.append(f"PDF 页数较多({page_count})，仅提取前 20 页")
        for page in pdf.pages[:20]:
            try:
                text = page.extract_text() or ""
            except Exception:  # noqa: BLE001
                warnings.append("单页 extract_text 失败，已跳过该页")
                continue
            if text.strip():
                chunks.append(text.strip())
    text = "\n".join(chunks)
    if len(text.strip()) < 40:
        warnings.append("PDF 文本很少，可能是扫描件/图片版。")
        if enable_ocr:
            if ocr_engine != "paddleocr":
                warnings.append(f"不支持的 OCR 引擎: {ocr_engine}")
            else:
                try:
                    # 与平台主链共用同一套 OCR 兜底；扫描图只在内存中处理，
                    # 评测产物仍然只写脱敏文本，不写回图片 PDF。
                    from app.services.parser import _parse_pdf

                    ocr_text, ocr_warnings = _parse_pdf(path)
                    warnings.extend(ocr_warnings)
                    if ocr_text.strip():
                        text = ocr_text
                    else:
                        warnings.append("OCR 未提取到有效正文。")
                except Exception as exc:  # noqa: BLE001 - 单文件失败不阻断批处理
                    warnings.append(f"OCR 失败: {type(exc).__name__}")
        else:
            warnings.append("未启用 OCR，扫描件仅输出不可读状态。")
    return text, warnings


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def ensure_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def copy_sidecar(src: Path, dst: Path) -> None:
    ensure_dir(dst.parent)
    shutil.copy2(src, dst)


def chunked(items: Iterable[Any], size: int) -> list[list[Any]]:
    buf: list[Any] = []
    out: list[list[Any]] = []
    for item in items:
        buf.append(item)
        if len(buf) >= size:
            out.append(buf)
            buf = []
    if buf:
        out.append(buf)
    return out
