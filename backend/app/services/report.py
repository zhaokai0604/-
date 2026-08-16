from datetime import datetime
from pathlib import Path
from typing import Any

from app.core.config import settings
from app.services.scoring import label_for_score


def export_stem(filename: str, fallback: str = "简历") -> str:
    """生成用于下载文件名的安全短名称，不影响数据库中的原始文件名。"""
    import re

    stem = Path(str(filename or "")).stem.strip()
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", stem)
    stem = re.sub(r"\s+", " ", stem).strip(" .")
    return (stem or fallback)[:48]


def report_download_name(filename: str, record_id: int, fmt: str, version_no: int = 1) -> str:
    return f"简析智评_{export_stem(filename)}_v{int(version_no or 1)}_分析报告_{record_id}.{fmt}"


def rewrite_download_name(filename: str, record_id: int, version_no: int = 1) -> str:
    return f"简析智评_{export_stem(filename)}_v{int(version_no or 1)}_优化稿_{record_id}.docx"


def generate_reports(
    record_id: int,
    filename: str,
    result: dict[str, Any],
    *,
    formats: list[str] | None = None,
) -> dict[str, Path]:
    settings.reports_dir.mkdir(parents=True, exist_ok=True)
    stem = f"analysis_{record_id}"
    selected = formats or ["docx", "pdf"]
    paths: dict[str, Path] = {}
    if "docx" in selected:
        docx_path = settings.reports_dir / f"{stem}.docx"
        generate_docx_report(docx_path, filename, result)
        paths["docx"] = docx_path
    if "pdf" in selected:
        pdf_path = settings.reports_dir / f"{stem}.pdf"
        generate_pdf_report(pdf_path, filename, result)
        paths["pdf"] = pdf_path
    return paths


def record_to_report_payload(record: Any) -> dict[str, Any]:
    from app.api.serializers import public_analysis_mode_label
    from app.utils.json_tools import loads

    sections = loads(record.sections_json, {})
    if not isinstance(sections, dict):
        sections = {}
    match_result = loads(record.match_result_json, {})
    if not isinstance(match_result, dict):
        match_result = {}
    public_sections = {
        key: value for key, value in sections.items()
        if not str(key).startswith("_") and isinstance(value, list) and value
    }
    return {
        "record_id": record.id,
        "version_no": int(record.version_no or 1),
        "original_filename": record.original_filename,
        "target_position": record.target_position or "",
        "job_description_present": bool((record.job_description or "").strip()),
        "total_score": record.total_score,
        "scores": loads(record.scores_json, {}),
        "diagnosis": loads(record.diagnosis_json, []),
        "suggestions": loads(record.suggestions_json, []),
        "structured_suggestions": sections.get("_structured_suggestions", []),
        "match_result": match_result,
        "job_market_match": sections.get("_job_market_match", {}),
        "analysis_mode": record.analysis_mode,
        "analysis_mode_label": public_analysis_mode_label(record.analysis_mode, sections),
        "ai_fallback_reason": "",
        "parse_quality": sections.get("_parse_quality", "medium"),
        "parse_warnings": sections.get("_parse_warnings", []),
        "quality_warnings": sections.get("_quality_warnings", []),
        "score_reliability": sections.get("_score_reliability", "normal"),
        "evidence_coverage": sections.get("_evidence_coverage"),
        "layout_complexity": sections.get("_layout_complexity"),
        "sections": public_sections,
        "structured_profile": sections.get("_structured", {}),
        "missing_sections": sections.get("_missing_sections", []),
        "action_roadmap": sections.get("_action_roadmap", []),
    }


def ensure_report_file(record: Any, fmt: str) -> Path:
    """按需生成报告文件（PDF 默认延迟到下载时生成）。"""
    settings.reports_dir.mkdir(parents=True, exist_ok=True)
    stem = f"analysis_{record.id}"
    path = settings.reports_dir / f"{stem}.{fmt}"
    if path.exists():
        return path
    payload = record_to_report_payload(record)
    if fmt == "docx":
        generate_docx_report(path, record.original_filename, payload)
    elif fmt == "pdf":
        generate_pdf_report(path, record.original_filename, payload)
    else:
        raise ValueError(f"unsupported report format: {fmt}")
    return path


def generate_docx_report(path: Path, filename: str, result: dict[str, Any]) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("简历评价报告", level=0)
    doc.add_paragraph(f"分析记录：#{result.get('record_id', '')} · v{result.get('version_no', 1)}")
    doc.add_paragraph(f"原始文件：{filename}")
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"分析模式：{result.get('analysis_mode_label') or _analysis_mode_label(result.get('analysis_mode', 'offline'))}")
    _add_scope_docx(doc, result)
    _add_parse_quality_docx(doc, result)
    doc.add_heading("总评分", level=1)
    total = result.get("total_score", 0)
    doc.add_paragraph(f"{total} / 100（{_score_grade_label(total)}）")
    match = result.get("match_result", {})
    if match.get("match_rate") is not None:
        doc.add_paragraph(f"岗位关键词覆盖率：{match.get('match_rate')}%")
    doc.add_heading("分项评分", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "维度"
    table.rows[0].cells[1].text = "分数"
    for key, score in result.get("scores", {}).items():
        row = table.add_row().cells
        row[0].text = label_for_score(key)
        row[1].text = str(score)
    _add_list_section(doc, "问题诊断", result.get("diagnosis", []))
    _add_list_section(doc, "修改建议", result.get("suggestions", []))
    _add_sections_docx(doc, result)
    structured = result.get("structured_suggestions") or []
    if structured:
        doc.add_heading("依据材料补强建议", level=1)
        for item in structured[:8]:
            if not isinstance(item, dict):
                continue
            doc.add_paragraph(f"问题：{item.get('problem', '')}")
            doc.add_paragraph(f"证据：{item.get('evidence', '')}")
            doc.add_paragraph(f"方向：{item.get('direction', '')}")
            doc.add_paragraph(f"示例：{item.get('example', '')}")
    match = result.get("match_result", {})
    doc.add_heading("岗位匹配分析", level=1)
    doc.add_paragraph(f"采用岗位：{match.get('target_position') or '未识别，使用通用建议'}")
    doc.add_paragraph(f"岗位来源：{_target_source_label(match.get('target_source', 'generic'))}")
    doc.add_paragraph(match.get("summary", ""))
    if match.get("missing_keywords"):
        doc.add_paragraph("建议补充关键词：" + "、".join(match["missing_keywords"]))
    _add_job_market_basis_docx(doc, result.get("job_market_match") or {})
    _add_limits_docx(doc, result)
    doc.save(path)


def _add_scope_docx(doc: Any, result: dict[str, Any]) -> None:
    target = str(result.get("target_position") or "").strip()
    doc.add_heading("本次分析范围", level=1)
    if target:
        doc.add_paragraph(f"评价对象：面向“{target}”岗位的简历。")
    else:
        doc.add_paragraph("评价对象：未提供目标岗位，以下为通用简历质量分析，不代表任何具体岗位的录用判断。")
    doc.add_paragraph(f"岗位输入：{'已提供岗位或岗位要求' if result.get('job_description_present') or target else '未提供'}")
    doc.add_paragraph("报告用途：辅助修改和复核，不替代招聘方的人工判断。")


def _add_sections_docx(doc: Any, result: dict[str, Any]) -> None:
    sections = result.get("sections") or {}
    if not isinstance(sections, dict) or not sections:
        return
    labels = {"basic_info": "基本信息", "education": "教育经历", "internship": "实习经历", "projects": "项目经历", "campus": "校园实践", "skills": "技能特长", "awards": "荣誉奖项", "summary": "求职意向/自我评价", "others": "其他内容"}
    doc.add_heading("解析到的简历板块", level=1)
    for key, lines in sections.items():
        if not isinstance(lines, list) or not lines:
            continue
        doc.add_heading(labels.get(key, key), level=2)
        for line in lines[:30]:
            doc.add_paragraph(str(line), style="List Bullet")


def _add_limits_docx(doc: Any, result: dict[str, Any]) -> None:
    doc.add_heading("使用边界与复核建议", level=1)
    missing = result.get("missing_sections") or []
    if missing:
        doc.add_paragraph("未识别或未发现的板块：" + "、".join(str(item) for item in missing))
    for item in (result.get("action_roadmap") or [])[:5]:
        if isinstance(item, dict):
            doc.add_paragraph(str(item.get("title") or item.get("action") or item), style="List Bullet")
        else:
            doc.add_paragraph(str(item), style="List Bullet")
    doc.add_paragraph("请优先核验：姓名与联系方式、教育和经历时间、项目数据、岗位要求及系统标出的低置信度内容。")


def _add_parse_quality_docx(doc: Any, result: dict[str, Any]) -> None:
    quality = result.get("parse_quality")
    if not quality:
        return
    doc.add_heading("分析可信度（分层说明）", level=1)
    doc.add_paragraph("链路：本次分析流程已完成（成功率只证明能跑通，不等于解析或打分准确）。")
    doc.add_paragraph(f"可读：{_parse_readable_label(quality)}")
    doc.add_paragraph(f"可信：{_parse_quality_label(quality)}")
    if result.get("evidence_coverage") is not None:
        doc.add_paragraph(f"材料覆盖率：{float(result.get('evidence_coverage') or 0) * 100:.0f}%")
    if result.get("layout_complexity") is not None:
        doc.add_paragraph(f"版式复杂度：{float(result.get('layout_complexity') or 0):.1f}（复杂度本身不直接扣分）")
    if result.get("score_reliability") == "low_parse_capped":
        doc.add_paragraph("因正文提取不足，总分已做保守封顶处理——证据不足时不给出虚高结论。")
    warnings = result.get("parse_warnings") or []
    warnings = list(dict.fromkeys([*warnings, *(result.get("quality_warnings") or [])]))
    for item in warnings:
        doc.add_paragraph(item, style="List Bullet")


def _add_job_market_basis_docx(doc: Any, market: dict[str, Any]) -> None:
    basis = market.get("requirement_basis") if isinstance(market, dict) else None
    related = market.get("related_jobs") if isinstance(market, dict) else None
    if not isinstance(basis, dict) and not related:
        return
    doc.add_heading("岗位适配依据", level=1)
    if isinstance(basis, dict) and (basis.get("must_skills") or basis.get("education")):
        if basis.get("education"):
            doc.add_paragraph(f"学历要求：{basis.get('education')}")
        if basis.get("must_skills"):
            doc.add_paragraph("技能要求：" + " / ".join(basis.get("must_skills") or []))
        if basis.get("experience_requirements"):
            doc.add_paragraph("经历要求：" + "、".join(basis.get("experience_requirements") or []))
        doc.add_paragraph(
            f"当前简历证据：命中 {basis.get('hit_count', 0)} 项，缺失 {basis.get('miss_count', 0)} 项"
        )
        if basis.get("hit_skills"):
            doc.add_paragraph("已命中：" + "、".join(basis.get("hit_skills") or []))
        if basis.get("missing_skills"):
            doc.add_paragraph("仍缺失：" + "、".join(basis.get("missing_skills") or []))
    if isinstance(related, list) and related:
        doc.add_paragraph("相关岗位 Top5（本地已核验公开岗位库）：")
        for item in related[:5]:
            if not isinstance(item, dict):
                continue
            doc.add_paragraph(
                f"{item.get('rank', '-')}. {item.get('target_position', '')} · "
                f"{item.get('city', '')} · 匹配分 {item.get('match_percent', 0)}",
                style="List Bullet",
            )


def _add_list_section(doc: Any, title: str, items: list[str]) -> None:
    doc.add_heading(title, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def generate_pdf_report(path: Path, filename: str, result: dict[str, Any]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name = _register_chinese_font()
    styles = _build_pdf_styles(font_name)
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm)
    story: list[Any] = []
    story.append(Paragraph("简历评价报告", styles["Title"]))
    story.append(Paragraph(f"分析记录：#{result.get('record_id', '')} · v{result.get('version_no', 1)}", styles["Normal"]))
    story.append(Paragraph(f"原始文件：{filename}", styles["Normal"]))
    story.append(Paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles["Normal"]))
    story.append(Paragraph(f"分析模式：{result.get('analysis_mode_label') or _analysis_mode_label(result.get('analysis_mode', 'offline'))}", styles["Normal"]))
    _add_scope_pdf(story, styles, result)
    _add_parse_quality_pdf(story, styles, result)
    story.append(Spacer(1, 8))
    total = result.get("total_score", 0)
    story.append(Paragraph(f"总评分：{total} / 100（{_score_grade_label(total)}）", styles["Heading2"]))
    match = result.get("match_result", {})
    if match.get("match_rate") is not None:
        story.append(Paragraph(f"岗位关键词覆盖率：{match.get('match_rate')}%", styles["Normal"]))
    rows = [["维度", "分数"]] + [[label_for_score(key), str(score)] for key, score in result.get("scores", {}).items()]
    table = Table(rows, colWidths=[95 * mm, 40 * mm])
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font_name), ("GRID", (0, 0), (-1, -1), 0.5, "#B8C1CC"), ("BACKGROUND", (0, 0), (-1, 0), "#E8F0F7")]))
    story.append(table)
    _pdf_section(story, styles, "问题诊断", result.get("diagnosis", []))
    _pdf_section(story, styles, "修改建议", result.get("suggestions", []))
    _add_sections_pdf(story, styles, result)
    match = result.get("match_result", {})
    story.append(Paragraph("岗位匹配分析", styles["Heading2"]))
    story.append(Paragraph(f"采用岗位：{match.get('target_position') or '未识别，使用通用建议'}", styles["Normal"]))
    story.append(Paragraph(f"岗位来源：{_target_source_label(match.get('target_source', 'generic'))}", styles["Normal"]))
    story.append(Paragraph(match.get("summary", ""), styles["Normal"]))
    if match.get("missing_keywords"):
        story.append(Paragraph("建议补充关键词：" + "、".join(match["missing_keywords"]), styles["Normal"]))
    _add_job_market_basis_pdf(story, styles, result.get("job_market_match") or {})
    _add_limits_pdf(story, styles, result)
    doc.build(story)


def _add_scope_pdf(story: list[Any], styles: Any, result: dict[str, Any]) -> None:
    from reportlab.platypus import Paragraph

    story.append(Paragraph("本次分析范围", styles["Heading2"]))
    target = str(result.get("target_position") or "").strip()
    text = f"评价对象：面向“{target}”岗位的简历。" if target else "评价对象：未提供目标岗位，以下为通用简历质量分析，不代表任何具体岗位的录用判断。"
    story.append(Paragraph(text, styles["Normal"]))
    story.append(Paragraph(f"岗位输入：{'已提供岗位或岗位要求' if result.get('job_description_present') or target else '未提供'}", styles["Normal"]))
    story.append(Paragraph("报告用途：辅助修改和复核，不替代招聘方的人工判断。", styles["Normal"]))


def _add_sections_pdf(story: list[Any], styles: Any, result: dict[str, Any]) -> None:
    from reportlab.platypus import Paragraph

    sections = result.get("sections") or {}
    if not isinstance(sections, dict) or not sections:
        return
    labels = {"basic_info": "基本信息", "education": "教育经历", "internship": "实习经历", "projects": "项目经历", "campus": "校园实践", "skills": "技能特长", "awards": "荣誉奖项", "summary": "求职意向/自我评价", "others": "其他内容"}
    story.append(Paragraph("解析到的简历板块", styles["Heading2"]))
    for key, lines in sections.items():
        if not isinstance(lines, list) or not lines:
            continue
        story.append(Paragraph(labels.get(key, key), styles["Heading3"]))
        for line in lines[:30]:
            story.append(Paragraph("- " + str(line), styles["Normal"]))


def _add_limits_pdf(story: list[Any], styles: Any, result: dict[str, Any]) -> None:
    from reportlab.platypus import Paragraph

    story.append(Paragraph("使用边界与复核建议", styles["Heading2"]))
    missing = result.get("missing_sections") or []
    if missing:
        story.append(Paragraph("未识别或未发现的板块：" + "、".join(str(item) for item in missing), styles["Normal"]))
    for item in (result.get("action_roadmap") or [])[:5]:
        text = item.get("title") or item.get("action") if isinstance(item, dict) else item
        story.append(Paragraph("- " + str(text), styles["Normal"]))
    story.append(Paragraph("请优先核验：姓名与联系方式、教育和经历时间、项目数据、岗位要求及系统标出的低置信度内容。", styles["Normal"]))


def _build_pdf_styles(font_name: str) -> Any:
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font_name
        if hasattr(style, "leading") and hasattr(style, "fontSize"):
            style.leading = max(style.leading, style.fontSize + 4)
    return styles


def _add_parse_quality_pdf(story: list[Any], styles: Any, result: dict[str, Any]) -> None:
    from reportlab.platypus import Paragraph, Spacer

    quality = result.get("parse_quality")
    if not quality:
        return
    story.append(Spacer(1, 8))
    story.append(Paragraph("分析可信度（分层说明）", styles["Heading2"]))
    story.append(Paragraph("链路：本次分析流程已完成（成功率只证明能跑通，不等于解析或打分准确）。", styles["Normal"]))
    story.append(Paragraph(f"可读：{_parse_readable_label(quality)}", styles["Normal"]))
    story.append(Paragraph(f"可信：{_parse_quality_label(quality)}", styles["Normal"]))
    if result.get("evidence_coverage") is not None:
        story.append(Paragraph(f"材料覆盖率：{float(result.get('evidence_coverage') or 0) * 100:.0f}%", styles["Normal"]))
    if result.get("layout_complexity") is not None:
        story.append(Paragraph(f"版式复杂度：{float(result.get('layout_complexity') or 0):.1f}（复杂度本身不直接扣分）", styles["Normal"]))
    if result.get("score_reliability") == "low_parse_capped":
        story.append(Paragraph("因正文提取不足，总分已做保守封顶处理——证据不足时不给出虚高结论。", styles["Normal"]))
    warnings = list(dict.fromkeys([*(result.get("parse_warnings") or []), *(result.get("quality_warnings") or [])]))
    for item in warnings:
        story.append(Paragraph("- " + item, styles["Normal"]))


def _add_job_market_basis_pdf(story: list[Any], styles: Any, market: dict[str, Any]) -> None:
    from reportlab.platypus import Paragraph, Spacer

    basis = market.get("requirement_basis") if isinstance(market, dict) else None
    related = market.get("related_jobs") if isinstance(market, dict) else None
    if not isinstance(basis, dict) and not related:
        return
    story.append(Spacer(1, 8))
    story.append(Paragraph("岗位适配依据", styles["Heading2"]))
    if isinstance(basis, dict) and (basis.get("must_skills") or basis.get("education")):
        if basis.get("education"):
            story.append(Paragraph(f"学历要求：{basis.get('education')}", styles["Normal"]))
        if basis.get("must_skills"):
            story.append(Paragraph("技能要求：" + " / ".join(basis.get("must_skills") or []), styles["Normal"]))
        if basis.get("experience_requirements"):
            story.append(Paragraph("经历要求：" + "、".join(basis.get("experience_requirements") or []), styles["Normal"]))
        story.append(
            Paragraph(
                f"当前简历证据：命中 {basis.get('hit_count', 0)} 项，缺失 {basis.get('miss_count', 0)} 项",
                styles["Normal"],
            )
        )
    if isinstance(related, list) and related:
        story.append(Paragraph("相关岗位 Top5：", styles["Normal"]))
        for item in related[:5]:
            if not isinstance(item, dict):
                continue
            story.append(
                Paragraph(
                    f"- {item.get('rank', '-')}. {item.get('target_position', '')} · "
                    f"{item.get('city', '')} · 匹配分 {item.get('match_percent', 0)}",
                    styles["Normal"],
                )
            )


def _pdf_section(story: list[Any], styles: Any, title: str, items: list[str]) -> None:
    from reportlab.platypus import Paragraph, Spacer

    story.append(Spacer(1, 8))
    story.append(Paragraph(title, styles["Heading2"]))
    for item in items:
        story.append(Paragraph("- " + item, styles["Normal"]))


_FONT_NAME: str | None = None


def _register_chinese_font() -> str:
    global _FONT_NAME
    if _FONT_NAME:
        return _FONT_NAME
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = list(settings.fonts_dir.glob("*.ttf")) + list(settings.fonts_dir.glob("*.otf"))
    windows_fonts = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]
    linux_fonts = [
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc"),
    ]
    candidates.extend(path for path in windows_fonts if path.exists())
    candidates.extend(path for path in linux_fonts if path.exists())
    for candidate in candidates:
        try:
            pdfmetrics.registerFont(TTFont("ResumeChinese", str(candidate)))
            _FONT_NAME = "ResumeChinese"
            return _FONT_NAME
        except Exception:
            continue
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        _FONT_NAME = "STSong-Light"
        return _FONT_NAME
    except Exception:
        pass
    _FONT_NAME = "Helvetica"
    return _FONT_NAME


def _analysis_mode_label(mode: str) -> str:
    labels = {
        "core": "规则分析",
        "deepseek": "增强分析已完成",
        "offline_fallback": "规则分析",
        "offline": "规则分析",
        "ai_first": "规则分析",
    }
    return labels.get(mode, mode)


def _target_source_label(source: str) -> str:
    labels = {
        "manual": "手动输入岗位",
        "selected": "点选推荐岗位",
        "detected": "简历识别岗位",
        "generic": "通用建议",
    }
    return labels.get(source, source)


def _parse_readable_label(quality: str) -> str:
    labels = {
        "high": "正文可读性较好，模块边界较清楚。",
        "medium": "正文可读，但版式或字段边界需人工核验。",
        "low": "正文可读性偏弱（扫描/碎片/噪声），可读≠板块识别准。",
    }
    return labels.get(quality, quality)


def _parse_quality_label(quality: str) -> str:
    labels = {
        "high": "分析可信度高：正文和核心模块识别较完整。",
        "medium": "分析可信度中等：该简历存在版式复杂或字段证据不足，结论可参考，建议优先补全文本信息。",
        "low": "分析可信度偏低：版式复杂/扫描质量较低/字段证据不足，建议优先补全可复制文本后再评估。",
    }
    return labels.get(quality, quality)


def _score_grade_label(score: Any) -> str:
    value = float(score or 0)
    if value >= 90:
        return "优秀"
    if value >= 80:
        return "良好"
    if value >= 70:
        return "中等"
    if value >= 60:
        return "待提升"
    return "需加强"


def generate_rewrite_report(record_id: int, filename: str, rewrite_preview: dict[str, Any]) -> Path:
    """导出基于初稿的完整优化稿；若无 optimized_resume 则回退为逐条改写参考。"""
    from docx import Document

    settings.reports_dir.mkdir(parents=True, exist_ok=True)
    path = settings.reports_dir / f"rewrite_{record_id}.docx"
    doc = Document()
    optimized = rewrite_preview.get("optimized_resume") if isinstance(rewrite_preview, dict) else None
    if isinstance(optimized, dict) and (optimized.get("sections") or optimized.get("document_text")):
        doc.add_heading(str(optimized.get("title") or "简历优化稿"), level=0)
        doc.add_paragraph(f"来源初稿：{filename}")
        doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(str(optimized.get("summary") or rewrite_preview.get("summary") or ""))
        target = optimized.get("target_position") or rewrite_preview.get("target_position") or ""
        if target:
            doc.add_paragraph(f"目标岗位：{target}")
        doc.add_paragraph("说明：本文档在初稿基础上打磨生成；【待补充】处请填入真实数据，系统未编造经历。")

        section_labels = {
            "basic_info": "基本信息",
            "education": "教育经历",
            "internship": "实习经历",
            "projects": "项目经历",
            "campus": "校园实践",
            "skills": "技能特长",
            "awards": "荣誉奖项",
            "summary": "自我评价/求职意向",
            "others": "其他",
        }
        sections = optimized.get("sections") if isinstance(optimized.get("sections"), dict) else {}
        ordered = [
            "basic_info",
            "education",
            "internship",
            "projects",
            "campus",
            "skills",
            "awards",
            "summary",
            "others",
        ]
        for key in ordered:
            lines = sections.get(key) or []
            if not lines:
                continue
            doc.add_heading(section_labels.get(key, key), level=1)
            for line in lines:
                doc.add_paragraph(str(line))
        for key, lines in sections.items():
            if key in ordered or not lines:
                continue
            doc.add_heading(section_labels.get(key, key), level=1)
            for line in lines:
                doc.add_paragraph(str(line))

        diffs = optimized.get("diffs") if isinstance(optimized.get("diffs"), list) else []
        if diffs:
            doc.add_heading("改动对照（初稿 → 优化稿）", level=1)
            for item in diffs[:20]:
                if not isinstance(item, dict):
                    continue
                doc.add_heading(str(item.get("section") or "改动"), level=2)
                if item.get("original"):
                    doc.add_paragraph(f"初稿：{item.get('original')}")
                doc.add_paragraph(f"优化稿：{item.get('optimized')}")
        doc.save(path)
        return path

    doc.add_heading("简历优化表达参考", level=0)
    doc.add_paragraph(f"来源文件：{filename}")
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(rewrite_preview.get("summary", ""))
    doc.add_paragraph(f"目标岗位：{rewrite_preview.get('target_position', '')}")
    for item in rewrite_preview.get("items") or []:
        if not isinstance(item, dict):
            continue
        doc.add_heading(str(item.get("section", "优化项")), level=2)
        doc.add_paragraph(f"原文：{item.get('original', '')}")
        doc.add_paragraph(f"参考改写：{item.get('suggested', '')}")
        doc.add_paragraph(f"优化重点：{item.get('focus', '')}")
    doc.save(path)
    return path
